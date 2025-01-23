import psycopg2
from psycopg2.extras import RealDictCursor, DictCursor
from datetime import date, datetime
import textwrap
import openpyxl

class PostgresConnector:
    def __init__(self, connection) -> None:
        """
        Initializes the PostgresConnector with a connection and sets the cursor to use RealDictCursor.

        Args:
            connection (dict): A dictionary containing the connection parameters for psycopg2.
        """
        self.conn = psycopg2.connect(**connection)
        self.cursor = self.conn.cursor(cursor_factory=RealDictCursor)  
        self._set_timezone_utc()

    def _set_timezone_utc(self):
        """
        Set the timezone for the database connection to UTC.
        """
        if self.conn:
            self.cursor.execute("SET TIME ZONE 'UTC';")
            
            
    def execute_cr_table_query(self, query):
        """
        Execute a single query.
        """
        try:
            if self.cursor:
                self.cursor.execute(query)
                self.conn.commit()
                print("Query executed successfully.")
            else:
                print("No active database connection.")
        except psycopg2.Error as e:
            print("Error executing query:", e)

    def execute_query(self, query, params=None, print_query: bool = True):
        """
        Execute a SQL query and return the results as a list of dictionaries where each dictionary
        represents a row with column names as keys.

        Args:
            query (str): The SQL query string to be executed.
            params (tuple, optional): parameters to pass with the query.
            print_query (bool, optional): For printing query on the terminal.

        Returns:
            list: A list of dictionaries, where each dictionary represents a row with column names as keys,
                  if the query is a SELECT statement.
            int: The number of affected rows for non-SELECT queries.
        """
        try:
            if print_query:
                print("\n", textwrap.dedent(query))
            self.cursor.execute(query, params)
            if query.strip().upper().startswith(("SELECT", "WITH")):
                result = self.cursor.fetchall()
                result_list = [dict(row) for row in result]
                return result_list
            else:
                affected_rows = self.cursor.rowcount
                self.conn.commit()
                return affected_rows
        except Exception as e:
            self.conn.rollback()
            print("\n", textwrap.dedent(query))
            raise e

    def close(self):
        """
        Closes the database connection and cursor.
        """
        self.cursor.close()
        self.conn.close()

    def copy_from_csv(self, csv_buffer, schema_name, table_name) -> None:
        """
        Copy data from a CSV buffer to a database table.

        Args:
            csv_buffer (Any): A file-like object containing CSV data.
            schema_name (str): The name of the schema containing the table.
            table_name (str): The name of the table to copy data into.

        Raises:
            psycopg2.Error: If there's an error during the copy operation.
        """
        copy_query = (
            f"COPY {schema_name}.{table_name} FROM STDIN WITH CSV HEADER DELIMITER ',' NULL AS ''"
        )
        try:
            self.cursor.copy_expert(copy_query, csv_buffer)
            self.conn.commit()
        except psycopg2.Error as e:
            self.conn.rollback()
            raise e
        
        
    def insert_from_excel(self, schema_name, table_name, file_path):
        """
        Create a table and insert data from an Excel file into the specified schema and table.
        """
        try:
            if not self.cursor:
                print("No active database connection.")
                return

            # Load Excel file
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook.active

            # Extract headers and create table columns
            columns = [cell.value for cell in sheet[1]]
            column_definitions = ', '.join([f'"{col}" TEXT' for col in columns])

            # Create schema if not exists
            self.execute_query(f"CREATE SCHEMA IF NOT EXISTS {schema_name};")

            # Create table if not exists
            create_table_query = f"""
                CREATE TABLE IF NOT EXISTS {schema_name}.{table_name} (
                    {column_definitions}
                );
            """
            self.execute_query(create_table_query)

            # Prepare insert query
            placeholders = ', '.join(['%s' for _ in columns])
            quoted_columns = ', '.join([f'"{col}"' for col in columns])
            insert_query = f"INSERT INTO {schema_name}.{table_name} ({quoted_columns}) VALUES ({placeholders});"

            # Insert data from Excel into the table
            for row in sheet.iter_rows(min_row=2, values_only=True):
                self.cursor.execute(insert_query, row)

            # Commit all changes
            self.conn.commit()
            print("Data inserted successfully from Excel.")

        except Exception as e:
            print("Error during Excel data insertion:", e)